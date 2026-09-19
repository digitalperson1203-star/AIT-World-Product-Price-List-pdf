"""Extract the supplied AIT workbook into the static catalog used by the webpage."""

import json
import re
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path(r"C:\Users\sorab\Downloads\product excel for project.xlsx")
ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets" / "product-images"
OUTPUT = ROOT / "catalog-data.js"


def clean(value):
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


def money(value):
    return value if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def slug(value):
    value = re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")
    return value or "collection"


def main():
    ASSETS.mkdir(parents=True, exist_ok=True)
    formulas = load_workbook(SOURCE, data_only=False)
    cached = load_workbook(SOURCE, data_only=True)

    collections = []
    terms = []
    exported_images = 0
    fallback_images = {}

    for source_sheet in formulas.worksheets:
        if source_sheet.sheet_state != "visible":
            continue

        values_sheet = cached[source_sheet.title]
        collection_slug = slug(source_sheet.title)
        image_by_row = {}
        for image in source_sheet._images:
            anchor = image.anchor
            row = anchor._from.row + 1
            col = anchor._from.col + 1
            if col == 2 and row not in image_by_row:
                image_by_row[row] = image

        products = []
        for row in range(10, source_sheet.max_row + 1):
            number = values_sheet.cell(row, 1).value
            description = clean(values_sheet.cell(row, 3).value)
            if not isinstance(number, (int, float)) or not description:
                continue

            base = money(values_sheet.cell(row, 5).value)
            discount = money(values_sheet.cell(row, 6).value)
            discounted = money(values_sheet.cell(row, 7).value)
            if base is not None and discounted is None and discount is not None:
                discounted = round(base * (100 - discount) / 100, 2)

            extra_prices = []
            for col in (8, 9):
                candidate = money(values_sheet.cell(row, col).value)
                if candidate is not None:
                    extra_prices.append(candidate)

            image_path = None
            if row in image_by_row:
                image = image_by_row[row]
                extension = image.format.lower() if image.format else "png"
                filename = f"{collection_slug}-row-{row}.{extension}"
                (ASSETS / filename).write_bytes(image._data())
                image_path = f"assets/product-images/{filename}"
                exported_images += 1
                fallback_images.setdefault(description, image_path)
            elif description in fallback_images:
                source_image = ROOT / fallback_images[description]
                extension = source_image.suffix
                filename = f"{collection_slug}-row-{row}{extension}"
                destination = ASSETS / filename
                destination.write_bytes(source_image.read_bytes())
                image_path = f"assets/product-images/{filename}"
                exported_images += 1

            products.append(
                {
                    "id": f"{collection_slug}-row-{row}",
                    "number": int(number),
                    "name": description,
                    "module": clean(values_sheet.cell(row, 4).value),
                    "unitPrice": base,
                    "discountPercent": discount,
                    "discountedPrice": discounted,
                    "sourcePrices": extra_prices,
                    "image": image_path,
                }
            )

        for row in range(1, values_sheet.max_row + 1):
            if clean(values_sheet.cell(row, 2).value) == "Terms & Conditions:":
                for term_row in range(row + 1, values_sheet.max_row + 1):
                    term = clean(values_sheet.cell(term_row, 2).value)
                    if term and term not in terms:
                        terms.append(term)
                break

        collections.append(
            {
                "id": collection_slug,
                "name": source_sheet.title.strip(),
                "technology": clean(values_sheet.cell(8, 5).value),
                "products": products,
            }
        )

    payload = {"collections": collections, "terms": terms}
    OUTPUT.write_text("window.CATALOG_DATA = " + json.dumps(payload, ensure_ascii=False, indent=2) + ";\n", encoding="utf-8")
    print(f"Exported {sum(len(c['products']) for c in collections)} products and {exported_images} images.")
    for collection in collections:
        missing = sum(1 for p in collection["products"] if not p["image"])
        print(f"{collection['name']}: {len(collection['products'])} products, {missing} images missing")


if __name__ == "__main__":
    main()
